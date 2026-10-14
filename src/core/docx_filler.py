"""Word 文档回填模块（PRD 3.2.5 / 4.6）。

核心思路：在原 .docx 副本上操作，逐段落替换文本内容，保留所有格式属性。

段落替换策略（PRD 4.6 伪代码）：
- 单 run 段落：直接替换 run 文本，格式随 run 保留
- 多 run 段落（混合格式）：新文本放入第一个 run，其余 run 清空，
  使整段继承第一个 run 的字符格式
- 无 run 空段落：add_run 写入新文本
- 新增段落：深拷贝锚点段落的 XML（继承段落格式与字符格式），
  插入到锚点之后，再替换文本

错误处理（PRD 8）：段落匹配失败时保留原段落不动，记录日志，继续处理其他段落。
"""

from __future__ import annotations

import copy
import logging
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from .docx_parser import iter_document_paragraphs

logger = logging.getLogger(__name__)

# Word 文件名中的非法字符（PRD 3.2.5：输出文件名含岗位关键词）
_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]')


@dataclass
class FillResult:
    """回填执行结果统计。"""

    replaced: int = 0
    skipped: int = 0
    inserted: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def sanitize_for_filename(text: str, max_len: int = 12) -> str:
    """清洗岗位关键词，使其可安全用于文件名。"""
    cleaned = _INVALID_FILENAME_CHARS.sub("", text).strip()
    # 去掉全部空格与常见分隔符，压缩过长的关键词
    cleaned = cleaned.replace(" ", "")[:max_len]
    return cleaned or "岗位"


def build_output_filename(
    original_path: str, job_keyword: str, date_str: str, avoid_dir: str | None = None
) -> str:
    """生成输出文件名：`原文件名_优化_岗位关键词_日期.docx`（PRD 3.2.5）。

    avoid_dir 非空时若同名文件已存在，自动追加 (2)、(3) 序号。
    """
    stem = Path(original_path).stem
    keyword = sanitize_for_filename(job_keyword)
    base = f"{stem}_优化_{keyword}_{date_str}"
    candidate = f"{base}.docx"
    if avoid_dir:
        seq = 2
        while (Path(avoid_dir) / candidate).exists():
            candidate = f"{base}({seq}).docx"
            seq += 1
    return candidate


# ---------------------------------------------------------------------------
# 段落文本替换（保留格式）
# ---------------------------------------------------------------------------

def _replace_paragraph_text(paragraph, new_text: str) -> None:
    """将段落文本替换为 new_text，保留段落与字符格式（PRD 4.6 策略）。"""
    if not new_text:
        # 明确清空段落（AI 计划中 action=delete / 置空）
        for run in paragraph.runs:
            run.text = ""
        return

    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return

    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


# ---------------------------------------------------------------------------
# 新增段落插入
# ---------------------------------------------------------------------------

def _insert_paragraph_after(anchor_paragraph, content: str):
    """在锚点段落后插入新段落，继承锚点的段落格式与字符格式。

    实现：深拷贝锚点段落的 <w:p> XML，插入其后，再替换文本。
    深拷贝保留了 pPr（段落属性）与 rPr（字符属性），视觉风格自然延续。
    """
    from docx.text.paragraph import Paragraph

    new_p_xml = copy.deepcopy(anchor_paragraph._p)
    anchor_paragraph._p.addnext(new_p_xml)
    new_paragraph = Paragraph(new_p_xml, anchor_paragraph._parent)
    _replace_paragraph_text(new_paragraph, content)
    return new_paragraph


# ---------------------------------------------------------------------------
# 回填主流程
# ---------------------------------------------------------------------------

def fill_resume(
    template_path: str,
    output_path: str,
    paragraphs_map: dict[str, str],
    new_paragraphs: list[dict] | None = None,
) -> FillResult:
    """将改写结果回填到原 Word 文档副本。

    Args:
        template_path: 原 Word 文件路径（不会被修改）
        output_path: 输出文件路径（先复制副本再修改）
        paragraphs_map: {paragraph_id: 新文本}；新文本为空表示清空该段
        new_paragraphs: [{"insert_after": "p5", "content": "..."}, ...]
                        新增段落，继承锚点段落格式

    Returns:
        FillResult: 替换/跳过/新增计数与错误信息
    """
    result = FillResult()
    new_paragraphs = new_paragraphs or []

    # 1) 复制副本（PRD 3.2.5：在原文档副本上修改，不覆盖原文件）
    output_dir = Path(output_path).parent
    output_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)

    doc = Document(output_path)

    # 2) 建立 ID -> 段落对象 映射（与解析器完全一致的顺序）
    id_to_paragraph = {}
    for para_id, para in iter_document_paragraphs(doc):
        id_to_paragraph[para_id] = para

    # 3) 逐段替换
    for para_id, new_text in paragraphs_map.items():
        para = id_to_paragraph.get(para_id)
        if para is None:
            # PRD 8：匹配失败保留原段落，记录，继续
            result.skipped += 1
            result.errors.append(f"段落 {para_id} 未找到，已保留原文")
            logger.warning("回填跳过：段落 %s 未找到", para_id)
            continue
        try:
            _replace_paragraph_text(para, new_text)
            result.replaced += 1
        except Exception as exc:  # 单段失败不影响其他段落
            result.skipped += 1
            result.errors.append(f"段落 {para_id} 替换失败：{exc}")
            logger.exception("回填失败：段落 %s", para_id)

    # 4) 插入新增段落（倒序插入以支持同一锚点后追加多条时顺序正确）
    for entry in reversed(new_paragraphs):
        anchor_id = entry.get("insert_after", "")
        content = entry.get("content", "")
        if not content or not str(content).strip():
            continue
        anchor = id_to_paragraph.get(anchor_id)
        if anchor is None:
            result.skipped += 1
            result.errors.append(
                f"新增段落锚点 {anchor_id} 未找到，该新增内容已跳过"
            )
            logger.warning("新增段落跳过：锚点 %s 未找到", anchor_id)
            continue
        try:
            _insert_paragraph_after(anchor, str(content))
            result.inserted += 1
        except Exception as exc:
            result.skipped += 1
            result.errors.append(f"新增段落插入失败（锚点 {anchor_id}）：{exc}")
            logger.exception("新增段落失败：锚点 %s", anchor_id)

    # 5) 保存
    doc.save(output_path)
    return result
