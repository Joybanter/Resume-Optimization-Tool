"""Word 文档解析模块（PRD 3.2.1 / 3.2.4 / 4.4）。

基于 python-docx 解析 .docx 文件，为每个段落分配稳定 ID：
- 正文段落：p0, p1, p2, ...（按文档顺序）
- 表格内段落：table_{t}_{r}_{c}_{k}
  （第 t 个表格、第 r 行、第 c 列单元格中的第 k 个段落）

段落 ID 是解析、AI 改写、回填三个环节的对齐锚点，
docx_parser 与 docx_filler 必须使用完全一致的遍历顺序。

异常处理（PRD 3.2.1 / 8）：
- DocFormatError：.doc 老格式 / 非 docx / 损坏文件
- EmptyResumeError：空文档或纯图片型简历
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject


class ResumeParseError(Exception):
    """简历解析相关错误基类。"""


class DocFormatError(ResumeParseError):
    """文件格式不正确（.doc 老格式 / 非 .docx / 损坏）。"""


class EmptyResumeError(ResumeParseError):
    """未检测到可识别的文本内容（空文档或图片型简历）。"""


@dataclass
class ParagraphInfo:
    """单个段落的解析结果。"""

    para_id: str
    text: str
    style_name: str = ""
    in_table: bool = False
    is_list_item: bool = False
    is_heading: bool = False
    bold: bool = False
    font_size_pt: float | None = None
    alignment: str = ""

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


@dataclass
class ParsedDocument:
    """一份简历文档的完整解析结果。"""

    source_path: str
    file_name: str = ""
    paragraphs: list[ParagraphInfo] = field(default_factory=list)
    table_count: int = 0
    image_count: int = 0
    parse_warnings: list[str] = field(default_factory=list)

    @property
    def total_chars(self) -> int:
        return sum(len(p.text) for p in self.paragraphs)

    @property
    def non_empty_count(self) -> int:
        return sum(1 for p in self.paragraphs if not p.is_empty)

    def get(self, para_id: str) -> ParagraphInfo | None:
        for p in self.paragraphs:
            if p.para_id == para_id:
                return p
        return None

    def to_prompt_text(self) -> str:
        """生成带段落 ID 前缀的全文，供 LLM 阅读与对齐。

        格式：[p0] 段落内容
        """
        lines = []
        for p in self.paragraphs:
            if p.is_empty and not p.in_table:
                # 空段落占位（保留编号连续性，便于 AI 理解结构）
                lines.append(f"[{p.para_id}]")
            else:
                lines.append(f"[{p.para_id}] {p.text}")
        return "\n".join(lines)

    def to_plain_text(self) -> str:
        return "\n".join(p.text for p in self.paragraphs if not p.is_empty)


# ---------------------------------------------------------------------------
# 魔数检测（区分 .doc 与 .docx / 损坏文件）
# ---------------------------------------------------------------------------

_DOCX_MAGIC = b"PK\x03\x04"      # docx 实为 zip 包
_DOC_MAGIC = b"\xd0\xcf\x11\xe0"  # 老版 .doc（OLE2 复合文档）


def _check_file_signature(path: str) -> None:
    """按文件头魔数判断真实格式。

    - 扩展名为 .doc：直接提示另存为 .docx（PRD 8）
    - 文件头为 OLE2：即使扩展名是 .docx 也按 .doc 处理
    - 非 zip：文件损坏或非 Word 文档
    """
    suffix = Path(path).suffix.lower()
    if suffix == ".doc":
        raise DocFormatError(
            "检测到旧版 .doc 格式，请在 Word 中将其另存为 .docx 后重新上传。"
        )
    if suffix != ".docx":
        raise DocFormatError("仅支持 .docx 格式的 Word 文档。")
    try:
        with open(path, "rb") as fh:
            head = fh.read(4)
    except OSError as exc:
        raise DocFormatError(f"无法读取文件：{exc}") from exc
    if head.startswith(_DOC_MAGIC):
        raise DocFormatError(
            "文件内容为旧版 .doc 格式，请在 Word 中将其另存为 .docx 后重新上传。"
        )
    if not head.startswith(_DOCX_MAGIC):
        raise DocFormatError("文件已损坏或不是有效的 Word 文档。")
    try:
        if not zipfile.is_zipfile(path):
            raise DocFormatError("文件已损坏或不是有效的 Word 文档。")
    except (OSError, zipfile.BadZipFile) as exc:
        if isinstance(exc, zipfile.BadZipFile):
            raise DocFormatError("文件已损坏或不是有效的 Word 文档。") from exc
        raise DocFormatError(f"无法读取文件：{exc}") from exc


# ---------------------------------------------------------------------------
# 解析主流程
# ---------------------------------------------------------------------------

def parse_docx(path: str) -> ParsedDocument:
    """解析 .docx 简历，返回 ParsedDocument。

    Raises:
        DocFormatError: 格式不正确
        EmptyResumeError: 无可识别文本
    """
    _check_file_signature(path)

    try:
        doc = Document(path)
    except Exception as exc:  # python-docx 对损坏文件抛多种异常
        raise DocFormatError(f"文档解析失败：{exc}") from exc

    parsed = ParsedDocument(source_path=str(path), file_name=Path(path).name)

    # 1) 正文段落
    for index, para in enumerate(doc.paragraphs):
        parsed.paragraphs.append(_make_paragraph_info(para, f"p{index}", in_table=False))

    # 2) 表格内段落
    parsed.table_count = len(doc.tables)
    for t, table in enumerate(doc.tables):
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                for k, para in enumerate(cell.paragraphs):
                    para_id = f"table_{t}_{r}_{c}_{k}"
                    parsed.paragraphs.append(
                        _make_paragraph_info(para, para_id, in_table=True)
                    )

    # 3) 图片统计（用于"图片型简历"判断）
    try:
        parsed.image_count = len(doc.inline_shapes)
    except Exception:  # pragma: no cover - inline_shapes 极少数情况解析异常
        parsed.image_count = 0

    # 4) 空文档 / 图片型简历检测（PRD 3.2.1）
    if parsed.total_chars < 10:
        if parsed.image_count > 0:
            raise EmptyResumeError(
                "未检测到可识别的文本内容：该文档疑似纯图片型简历，"
                "请使用可复制文字的 .docx 文档。"
            )
        raise EmptyResumeError(
            "未检测到可识别的文本内容：文档为空或无可读取的文字。"
        )

    return parsed


def _make_paragraph_info(para, para_id: str, in_table: bool) -> ParagraphInfo:
    """从 python-docx 段落对象提取文本与格式摘要。"""
    try:
        style_name = para.style.name if para.style is not None else ""
    except Exception:
        style_name = ""

    is_list_item = False
    try:
        if para._p.pPr is not None and para._p.pPr.numPr is not None:
            is_list_item = True
    except Exception:
        pass

    bold = False
    font_size = None
    for run in para.runs:
        if run.bold:
            bold = True
        if run.font.size is not None and font_size is None:
            font_size = run.font.size.pt
        if bold and font_size is not None:
            break

    alignment = ""
    try:
        align = para.alignment
        if align is not None:
            alignment = str(align).split(".")[-1].split(" ")[0] if align else ""
    except Exception:
        alignment = ""

    is_heading = bool(style_name) and (
        style_name.lower().startswith("heading") or "标题" in style_name
    )

    return ParagraphInfo(
        para_id=para_id,
        text=para.text,
        style_name=style_name,
        in_table=in_table,
        is_list_item=is_list_item,
        is_heading=is_heading,
        bold=bold,
        font_size_pt=font_size,
        alignment=alignment,
    )


# ---------------------------------------------------------------------------
# 供回填模块复用：以相同顺序枚举 Document 的段落对象
# ---------------------------------------------------------------------------

def iter_document_paragraphs(doc: DocumentObject):
    """按与 parse_docx 完全一致的顺序枚举 (para_id, paragraph对象)。

    回填模块通过本函数定位段落，保证 ID 对齐。
    """
    for index, para in enumerate(doc.paragraphs):
        yield f"p{index}", para
    for t, table in enumerate(doc.tables):
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                for k, para in enumerate(cell.paragraphs):
                    yield f"table_{t}_{r}_{c}_{k}", para


# ---------------------------------------------------------------------------
# Token 预估（PRD 3.2.2：JD 框下方显示字符数与预估 Token）
# ---------------------------------------------------------------------------

def estimate_tokens(text: str) -> int:
    """粗略估算 Token 数：中文约 1 字 ≈ 0.6 token，英文按 4 字符 ≈ 1 token。"""
    if not text:
        return 0
    cjk = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")
    other = len(text) - cjk
    return int(cjk * 0.6 + other * 0.25) + 1
